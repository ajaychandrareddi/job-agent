"""Extract and normalise text from PDF and DOCX resumes."""

from __future__ import annotations

import io
import re
import unicodedata

from job_agent.models.resume import ResumeFileType

_LIGATURES = {"ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl"}
_BULLETS = re.compile("^[ \t]*[•●▪◦‣⁃∙\\-*–—]+[ \t]*", re.M)
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class TextExtractionError(ValueError):
    pass


def extract_text(content: bytes, file_type: ResumeFileType) -> str:
    try:
        raw = _extract_pdf(content) if file_type is ResumeFileType.PDF else _extract_docx(content)
    except TextExtractionError:
        raise
    except Exception as exc:  # parser libraries raise many types on malformed input
        raise TextExtractionError(f"could not read {file_type.value} file") from exc
    text = normalize_text(raw)
    if len(text) < 50:
        raise TextExtractionError(
            "resume contains too little extractable text (scanned images are not supported)"
        )
    return text


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise TextExtractionError("password-protected PDFs are not supported")
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(parts)


def normalize_text(text: str) -> str:
    """NFKC-normalise, strip control characters and bullets, collapse blank runs."""
    for lig, repl in _LIGATURES.items():
        text = text.replace(lig, repl)
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = _CONTROL.sub("", text)
    text = _BULLETS.sub("", text)
    lines = [" ".join(line.split()) for line in text.split("\n")]
    out: list[str] = []
    for line in lines:
        if line or (out and out[-1]):
            out.append(line)
    return "\n".join(out).strip()
