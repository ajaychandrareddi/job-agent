from __future__ import annotations

import io
import os
from collections.abc import Iterator
from datetime import UTC, date, datetime, timedelta

import pytest
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session

from job_agent.models.job import RawPosting
from job_agent.models.resume import JobSearchPreferences
from job_agent.persistence.db import create_db_engine, create_session_factory
from job_agent.persistence.file_store import InMemoryResumeFileStore
from job_agent.persistence.orm import Base, User
from job_agent.repositories.resumes import UserRepository
from job_agent.services.resume_parser import ResumeParser
from job_agent.services.resume_service import ResumeService
from tests.db_guard import assert_test_database

TODAY = date(2026, 10, 1)

SAMPLE_RESUME_LINES = [
    "Jane Doe",
    "jane@example.com | (555) 123-4567",
    "Summary",
    "Backend engineer building data platforms in healthcare.",
    "Skills",
    "Languages: Python, Go, SQL",
    "Frameworks: FastAPI, Django",
    "Technologies",
    "AWS, Docker, Kubernetes, PostgreSQL",
    "Experience",
    "Senior Software Engineer at Acme Health, Jan 2020 - Present",
    "- Led migration of billing services to Kubernetes, cutting costs 30%",
    "Software Engineer at DataCorp, Jun 2016 - Dec 2019",
    "- Built ETL pipelines in Python",
    "Education",
    "B.S. Computer Science, State University, 2016",
    "Certifications",
    "AWS Certified Solutions Architect - Associate",
]


def make_docx(lines: list[str], *, table_rows: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for line in lines:
        document.add_paragraph(line)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_pdf(lines: list[str]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in lines:
        pdf.cell(0, 8, line.replace("–", "-"), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


@pytest.fixture
def sample_docx() -> bytes:
    return make_docx(SAMPLE_RESUME_LINES)


@pytest.fixture
def sample_pdf() -> bytes:
    return make_pdf(SAMPLE_RESUME_LINES)


@pytest.fixture
def engine() -> Iterator[Engine]:
    url = os.environ.get("TEST_DATABASE_URL_UNIT", "sqlite://")
    assert_test_database(url)
    eng = create_db_engine(url)
    Base.metadata.create_all(eng)
    yield eng
    Base.metadata.drop_all(eng)
    eng.dispose()


@pytest.fixture
def session(engine: Engine) -> Iterator[Session]:
    factory = create_session_factory(engine)
    s = factory()
    try:
        yield s
    finally:
        s.rollback()
        s.close()


@pytest.fixture
def user(session: Session) -> User:
    return UserRepository(session).get_or_create("auth0|alice", email="alice@example.com")


@pytest.fixture
def other_user(session: Session) -> User:
    return UserRepository(session).get_or_create("auth0|bob")


@pytest.fixture
def file_store() -> InMemoryResumeFileStore:
    return InMemoryResumeFileStore()


@pytest.fixture
def resume_service(session: Session, file_store: InMemoryResumeFileStore) -> ResumeService:
    return ResumeService(session, ResumeParser(), file_store)


@pytest.fixture
def approved_prefs() -> JobSearchPreferences:
    return JobSearchPreferences(
        preferred_titles=["Senior Software Engineer"], locations=["New York, NY"]
    )


# --------------------------------------------------------------------------- Phase 2

RUN_AT = datetime(2026, 10, 1, 21, 0, tzinfo=UTC)  # 5:00 PM America/New_York


def make_raw(external_id: str = "1", **overrides: object) -> RawPosting:
    """A complete, valid raw posting published two hours before RUN_AT."""
    data: dict[str, object] = {
        "external_id": external_id,
        "title": "Senior Software Engineer",
        "company_name": "Acme Health, Inc.",
        "location": "New York, NY",
        "employment_type": "Full-time",
        "description": "<p>Build APIs.</p><ul><li>5+ years of Python experience</li></ul>",
        "published_at": (RUN_AT - timedelta(hours=2)).isoformat(),
        "url": f"https://jobs.example.com/view/{external_id}?utm_source=feed",
        "apply_url": f"https://boards.greenhouse.io/acme/jobs/{external_id}",
        "is_open": True,
    }
    data.update(overrides)
    return RawPosting.model_validate(data)
